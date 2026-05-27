"""
ingest.py  —  Load, chunk, embed and store both source files into ChromaDB.

Sources:
  1. Emotions_vocabulary_final.xlsx  → one document per vocabulary word
  2. Task_1-_Neha.docx               → one document per section/paragraph

Run once (or re-run to refresh):
    python ingest.py
"""

import re
import uuid
from typing import List, Dict

import chromadb
from chromadb.utils import embedding_functions
import openpyxl
from docx import Document

from config import (
    CHROMA_DIR, COLLECTION_EMOTIONS, COLLECTION_TASK,
    EMBED_MODEL, XLSX_FILE, DOCX_FILE,
)
from ruler_mapper import ruler_summary, map_to_ruler


# ── Helpers ────────────────────────────────────────────────────────────────

def _chroma_client() -> chromadb.PersistentClient:
    return chromadb.PersistentClient(path=CHROMA_DIR)


def _embed_fn():
    return embedding_functions.SentenceTransformerEmbeddingFunction(
        model_name=EMBED_MODEL
    )


def _safe_str(val) -> str:
    """Convert any cell value to a clean string."""
    if val is None:
        return ""
    return str(val).strip()


# ── XLSX Ingestion ─────────────────────────────────────────────────────────

def load_emotions_xlsx(path: str) -> List[Dict]:
    """
    Parse the emotions vocabulary spreadsheet.
    Returns a list of dicts, one per vocabulary word.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    headers_raw = rows[0]

    # Normalise header names
    header_map = {
        "S.No.":            "sno",
        "Word":             "word",
        "Definition":       "definition",
        "Example":          "example",
        "Category":         "category",
        "Level":            "level",
        "Similar Words":    "similar_words",
        "Opposite Words":   "opposite_words",
        "Cultural Context": "cultural_context",
    }
    headers = [header_map.get(str(h).strip(), str(h).strip()) for h in headers_raw]

    records = []
    for row in rows[1:]:
        record = {headers[i]: _safe_str(v) for i, v in enumerate(row)}
        if not record.get("word"):        # skip empty rows
            continue
        records.append(record)

    print(f"[ingest] Loaded {len(records)} emotion words from XLSX.")
    return records


def ingest_emotions(records: List[Dict]) -> None:
    """Embed each emotion word and store in ChromaDB."""
    client = _chroma_client()
    ef     = _embed_fn()

    # Fresh collection each run
    try:
        client.delete_collection(COLLECTION_EMOTIONS)
    except Exception:
        pass
    col = client.create_collection(
        name=COLLECTION_EMOTIONS,
        embedding_function=ef,
        metadata={"hnsw:space": "cosine"},
    )

    docs, ids, metas = [], [], []
    for rec in records:
        doc_text = ruler_summary(rec)
        ruler    = map_to_ruler(rec)
        meta = {
            "source":           "emotions_vocabulary_xlsx",
            "word":             rec.get("word", ""),
            "category":         rec.get("category", ""),
            "level":            rec.get("level", ""),
            "definition":       rec.get("definition", ""),
            "example":          rec.get("example", ""),
            "similar_words":    rec.get("similar_words", ""),
            "opposite_words":   rec.get("opposite_words", ""),
            "cultural_context": rec.get("cultural_context", ""),
            "ruler_recognize":  ruler["Recognize"],
            "ruler_understand": ruler["Understand"],
            "ruler_label":      ruler["Label"],
            "ruler_express":    ruler["Express"],
            "ruler_regulate":   ruler["Regulate"],
        }
        docs.append(doc_text)
        ids.append(f"emotion_{rec.get('word','').lower().replace(' ','_')}_{uuid.uuid4().hex[:6]}")
        metas.append(meta)

    # Batch upsert (ChromaDB limit: 5000 per call)
    batch = 500
    for i in range(0, len(docs), batch):
        col.upsert(
            documents=docs[i:i+batch],
            ids=ids[i:i+batch],
            metadatas=metas[i:i+batch],
        )

    print(f"[ingest] Stored {len(docs)} emotion documents in '{COLLECTION_EMOTIONS}'.")


# ── DOCX Ingestion ─────────────────────────────────────────────────────────

def _split_docx_into_sections(path: str) -> List[Dict]:
    """
    Split the task document by headings into labelled sections.
    Each section carries the heading + body text.
    """
    doc = Document(path)
    sections = []
    current_heading = "Introduction"
    current_body: List[str] = []

    for para in doc.paragraphs:
        style = para.style.name if para.style else ""
        text  = para.text.strip()
        if not text:
            continue

        is_heading = (
            style.startswith("Heading")
            or style == "Title"
            or (len(text) < 120 and text.endswith((":", ")")))
        )

        if is_heading:
            if current_body:
                sections.append({
                    "heading": current_heading,
                    "body":    " ".join(current_body),
                })
                current_body = []
            current_heading = text
        else:
            current_body.append(text)

    # flush last section
    if current_body:
        sections.append({
            "heading": current_heading,
            "body":    " ".join(current_body),
        })

    print(f"[ingest] Split DOCX into {len(sections)} sections.")
    return sections


def ingest_task_doc(path: str) -> None:
    """Embed each document section and store in ChromaDB."""
    sections = _split_docx_into_sections(path)

    client = _chroma_client()
    ef     = _embed_fn()

    try:
        client.delete_collection(COLLECTION_TASK)
    except Exception:
        pass
    col = client.create_collection(
        name=COLLECTION_TASK,
        embedding_function=ef,
        metadata={"hnsw:space": "cosine"},
    )

    docs, ids, metas = [], [], []

    def _infer_source_type(heading: str) -> str:
        h = heading.lower()
        if "source type" in h or "voice journal" in h or "text entry" in h:
            return "source_type"
        if "journal entries" in h or "voice transcripts" in h:
            return "journal_entries"
        if "categorize entries" in h or "categorize" in h:
            return "categorization_rules"
        if "task 1" in h:
            return "task_1"
        return "task_docx"

    for i, sec in enumerate(sections):
        doc_text = f"Section: {sec['heading']}\n\n{sec['body']}"
        source_type = _infer_source_type(sec['heading'])
        meta = {
            "source":      "task_docx",
            "source_type": source_type,
            "heading":     sec["heading"],
            "section_index": str(i),
        }
        docs.append(doc_text)
        ids.append(f"task_sec_{i}_{uuid.uuid4().hex[:6]}")
        metas.append(meta)

    col.upsert(documents=docs, ids=ids, metadatas=metas)
    print(f"[ingest] Stored {len(docs)} task sections in '{COLLECTION_TASK}'.")


# ── Entry point ────────────────────────────────────────────────────────────

def run_ingestion() -> None:
    print("\n=== Starting ingestion ===")
    records = load_emotions_xlsx(XLSX_FILE)
    ingest_emotions(records)
    ingest_task_doc(DOCX_FILE)
    print("=== Ingestion complete ===\n")


if __name__ == "__main__":
    run_ingestion()
